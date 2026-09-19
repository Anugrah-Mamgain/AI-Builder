from __future__ import annotations

import io
from typing import Any

from pptx import Presentation as PptxPresentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from schemas import Presentation


def _as_dict(deck: Any) -> dict[str, Any]:
    if isinstance(deck, dict):
        return deck

    if hasattr(deck, "model_dump"):
        return deck.model_dump()

    if hasattr(deck, "dict"):
        return deck.dict()

    raise TypeError("deck must be a Presentation model or dictionary")


def _add_text(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    size: int,
    color: RGBColor,
    *,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )

    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True

    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0

    paragraph = frame.paragraphs[0]
    paragraph.text = str(text)
    paragraph.alignment = align

    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color

    return box


def _add_footer(slide, number: int, muted: RGBColor):
    _add_text(
        slide,
        "AI BUILDER",
        0.68,
        7.05,
        2.1,
        0.22,
        8,
        muted,
        bold=True,
    )

    _add_text(
        slide,
        f"{number:02d}",
        12.05,
        7.0,
        0.6,
        0.25,
        9,
        muted,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def export_pptx(deck: Presentation) -> bytes:
    """Convert a Presentation schema/model into PPTX bytes."""

    data = _as_dict(deck)

    presentation = PptxPresentation()

    # 16:9 widescreen
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    # Theme colors
    navy = RGBColor(15, 23, 42)
    indigo = RGBColor(79, 70, 229)
    violet = RGBColor(124, 58, 237)
    pink = RGBColor(219, 39, 119)
    white = RGBColor(255, 255, 255)
    ink = RGBColor(30, 41, 59)
    muted = RGBColor(100, 116, 139)
    pale = RGBColor(238, 242, 255)
    light_purple = RGBColor(196, 181, 253)
    light_gray = RGBColor(203, 213, 225)
    border = RGBColor(224, 231, 255)

    # -------------------------
    # Title slide
    # -------------------------
    title_slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )

    title_slide.background.fill.solid()
    title_slide.background.fill.fore_color.rgb = navy

    accent = title_slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0),
        Inches(0),
        Inches(0.18),
        Inches(7.5),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = pink
    accent.line.fill.background()

    _add_text(
        title_slide,
        "AI BUILDER",
        0.85,
        0.72,
        2.6,
        0.35,
        12,
        light_purple,
        bold=True,
    )

    _add_text(
        title_slide,
        data.get("title", "Generated Presentation"),
        0.85,
        2.0,
        11.2,
        1.4,
        34,
        white,
        bold=True,
    )

    _add_text(
        title_slide,
        f"Audience: {data.get('audience', 'General audience')}",
        0.88,
        4.0,
        9.9,
        0.6,
        18,
        light_gray,
    )

    _add_text(
        title_slide,
        "AI-GENERATED PRESENTATION",
        0.88,
        6.55,
        3.6,
        0.3,
        9,
        RGBColor(148, 163, 184),
        bold=True,
    )

    # -------------------------
    # Content slides
    # -------------------------
    for slide_number, slide_data in enumerate(
        data.get("slides", []),
        start=2,
    ):
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = white

        band = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(0.18),
            Inches(7.5),
        )

        band.fill.solid()
        band.fill.fore_color.rgb = (
            indigo if slide_number % 2 == 0 else violet
        )
        band.line.fill.background()

        _add_text(
            slide,
            slide_data.get("title", "Untitled Slide"),
            0.72,
            0.55,
            8.2,
            0.6,
            27,
            ink,
            bold=True,
        )

        _add_text(
            slide,
            "KEY POINTS",
            0.75,
            1.55,
            2,
            0.3,
            10,
            indigo,
            bold=True,
        )

        key_points = slide_data.get("key_points", [])[:5]

        for bullet_index, bullet in enumerate(key_points):
            y = 2.0 + bullet_index * 0.95

            dot = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(0.78),
                Inches(y + 0.08),
                Inches(0.18),
                Inches(0.18),
            )

            dot.fill.solid()
            dot.fill.fore_color.rgb = pink
            dot.line.fill.background()

            _add_text(
                slide,
                bullet,
                1.15,
                y,
                7.15,
                0.72,
                18,
                ink,
            )

        # Visual suggestion panel
        panel = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(9.0),
            Inches(1.48),
            Inches(3.55),
            Inches(4.95),
        )

        panel.fill.solid()
        panel.fill.fore_color.rgb = pale
        panel.line.color.rgb = border

        _add_text(
            slide,
            "VISUAL DIRECTION",
            9.38,
            1.9,
            2.8,
            0.3,
            10,
            violet,
            bold=True,
        )

        _add_text(
            slide,
            slide_data.get(
                "visual_suggestion",
                "Supporting visual",
            ),
            9.38,
            2.45,
            2.75,
            1.75,
            18,
            ink,
            bold=True,
        )

        _add_text(
            slide,
            "SPEAKER NOTES",
            9.38,
            4.8,
            2.5,
            0.3,
            10,
            muted,
            bold=True,
        )

        notes = slide_data.get("speaker_notes", "")

        _add_text(
            slide,
            notes,
            9.38,
            5.18,
            2.72,
            0.86,
            11,
            muted,
        )

        # Actual PowerPoint speaker notes
        slide.notes_slide.notes_text_frame.text = notes

        _add_footer(slide, slide_number, muted)

    # -------------------------
    # Audience questions slide
    # -------------------------
    questions = data.get("audience_questions", [])

    if questions:
        slide_number = len(presentation.slides) + 1

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = navy

        _add_text(
            slide,
            "Questions to Expect",
            0.8,
            0.65,
            8,
            0.65,
            28,
            white,
            bold=True,
        )

        for index, item in enumerate(questions[:4]):
            y = 1.65 + index * 1.25

            _add_text(
                slide,
                f"{index + 1:02d}",
                0.85,
                y,
                0.55,
                0.35,
                13,
                light_purple,
                bold=True,
            )

            _add_text(
                slide,
                item.get("question", "Question"),
                1.55,
                y - 0.03,
                10.5,
                0.4,
                17,
                white,
                bold=True,
            )

            _add_text(
                slide,
                item.get("suggested_answer", ""),
                1.55,
                y + 0.43,
                10.25,
                0.55,
                12,
                light_gray,
            )

        _add_footer(slide, slide_number, muted)

    # Serialize to bytes
    output = io.BytesIO()
    presentation.save(output)

    return output.getvalue()