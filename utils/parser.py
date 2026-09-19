from __future__ import annotations

import hashlib
import io
from pathlib import Path
from typing import Any

from config import ALLOWED_EXTENSIONS, MAX_UPLOAD_MB


MAX_FILE_SIZE_BYTES = MAX_UPLOAD_MB * 1024 * 1024


class ParserError(ValueError):
    """Base error for document parsing failures."""


class FileTooLargeError(ParserError):
    pass


class UnsupportedFileTypeError(ParserError):
    pass


class EmptyFileError(ParserError):
    pass


class NoExtractableTextError(ParserError):
    pass


def _read_input(uploaded_file: Any, filename: str | None = None) -> tuple[bytes, str]:
    """Accept Streamlit UploadedFile, bytes, or file-like objects."""
    if isinstance(uploaded_file, bytes):
        if not filename:
            raise UnsupportedFileTypeError(
                "Filename is required when passing raw bytes."
            )
        return uploaded_file, Path(filename).name

    if isinstance(uploaded_file, bytearray):
        if not filename:
            raise UnsupportedFileTypeError(
                "Filename is required when passing raw bytes."
            )
        return bytes(uploaded_file), Path(filename).name

    name = filename or getattr(uploaded_file, "name", None)
    if not name:
        raise UnsupportedFileTypeError("Uploaded file must have a filename.")

    # Streamlit UploadedFile and normal file-like objects
    if hasattr(uploaded_file, "getvalue"):
        data = uploaded_file.getvalue()
    elif hasattr(uploaded_file, "read"):
        data = uploaded_file.read()
    else:
        raise TypeError("Unsupported upload object.")

    if not isinstance(data, bytes):
        data = bytes(data)

    return data, Path(name).name


def _validate(data: bytes, filename: str) -> str:
    if not data:
        raise EmptyFileError("The uploaded file is empty.")

    if len(data) > MAX_FILE_SIZE_BYTES:
        raise FileTooLargeError(
            f"File exceeds the {MAX_UPLOAD_MB} MB upload limit."
        )

    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{extension or 'unknown'}'. "
            f"Allowed types: PDF, DOCX, TXT."
        )

    return extension


def _parse_txt(data: bytes) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("utf-8", errors="replace")

    text = text.strip()

    if not text:
        raise EmptyFileError("The text file contains no readable text.")

    return text


def _parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))

        chunks: list[str] = []

        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                chunks.append(page_text.strip())

        text = "\n".join(chunks).strip()

    except Exception as exc:
        raise ParserError(
            f"Unable to parse PDF: {type(exc).__name__}"
        ) from exc

    if not text:
        raise NoExtractableTextError(
            "PDF contains no extractable text. "
            "It may be a scanned or image-only PDF."
        )

    return text


def _parse_docx(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))

        chunks: list[str] = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Extract basic table text too.
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    chunks.append(" | ".join(cells))

        text = "\n".join(chunks).strip()

    except Exception as exc:
        raise ParserError(
            f"Unable to parse DOCX: {type(exc).__name__}"
        ) from exc

    if not text:
        raise EmptyFileError("DOCX contains no readable text.")

    return text


def extract_text(
    uploaded_file: Any,
    filename: str | None = None,
) -> str:
    """Extract text from PDF, DOCX or TXT.

    Raw extracted text is held only in memory and is returned to the
    preprocessing pipeline. This function never writes extracted text
    to disk, a database, logs, or telemetry.
    """
    data, name = _read_input(uploaded_file, filename)

    extension = _validate(data, name)

    if extension == ".txt":
        return _parse_txt(data)

    if extension == ".pdf":
        return _parse_pdf(data)

    if extension == ".docx":
        return _parse_docx(data)

    # Defensive guard; _validate() should make this unreachable.
    raise UnsupportedFileTypeError("Unsupported document type.")


def content_hash(
    uploaded_file: Any,
    filename: str | None = None,
) -> str:
    """Return SHA-256 for the original uploaded bytes."""
    data, _ = _read_input(uploaded_file, filename)
    return hashlib.sha256(data).hexdigest()