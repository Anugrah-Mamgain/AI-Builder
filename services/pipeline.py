"""P3 orchestration: parse -> redact -> generate -> save -> export.

Imports are lazy so the app remains demoable while P1 and P2 build their modules.
"""
from __future__ import annotations

import copy
import io
import logging
from pathlib import Path
from typing import Any
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

LOGGER = logging.getLogger(__name__)
_DEMO_HISTORY: list[dict[str, Any]] = []

MOCK_DECK = {
    "title": "AI in Indian Healthcare",
    "subtitle": "From operational efficiency to better patient outcomes",
    "slides": [
        {"title": "The Opportunity", "bullets": [
            "Growing patient demand is stretching hospital capacity",
            "Clinical teams lose time to repetitive administrative work",
            "Responsible AI augments decisions without replacing clinicians",
        ], "speaker_notes": "Lead with the human impact: more time for patients.",
         "visual_suggestion": "Before-and-after hospital workflow"},
        {"title": "Where AI Creates Value", "bullets": [
            "Automate appointment triage and routine documentation",
            "Predict capacity needs and optimize resources",
            "Surface relevant information at the point of care",
        ], "speaker_notes": "Tie every use case to a measurable KPI.",
         "visual_suggestion": "Patient, clinician, and hospital value map"},
        {"title": "Responsible Implementation", "bullets": [
            "Keep clinicians in the loop for consequential decisions",
            "Minimize and encrypt personally identifiable information",
            "Monitor accuracy, bias, latency, and adoption",
        ], "speaker_notes": "Governance is part of product design.",
         "visual_suggestion": "Circular governance diagram"},
        {"title": "90-Day Pilot Roadmap", "bullets": [
            "Weeks 1–3: select one workflow and baseline metrics",
            "Weeks 4–8: launch a controlled pilot with training",
            "Weeks 9–12: measure impact and decide whether to scale",
        ], "speaker_notes": "Close with a low-risk next step.",
         "visual_suggestion": "Horizontal roadmap with three milestones"},
    ],
    "qa": [
        {"question": "How is sensitive information protected?", "answer": "Use minimization, encryption, access controls, logs, and retention limits."},
        {"question": "How will we measure success?", "answer": "Compare time, quality, satisfaction, adoption, and cost against a baseline."},
        {"question": "Will AI replace staff?", "answer": "AI assists repetitive work; accountable humans retain final control."},
    ],
}


def _as_dict(model: Any) -> dict[str, Any]:
    if isinstance(model, dict):
        return model
    if hasattr(model, "model_dump"):
        return model.model_dump()
    if hasattr(model, "dict"):
        return model.dict()
    raise TypeError("Presentation must be a dictionary or Pydantic model.")


def _parse_document(data: bytes, name: str) -> tuple[str, bool]:
    try:
        from utils.parser import extract_text
    except (ImportError, ModuleNotFoundError):
        return (data.decode("utf-8", errors="replace") if name.lower().endswith(".txt")
                else f"Demo content extracted from {name}"), True
    try:
        try:
            return extract_text(data, name), False
        except TypeError:
            return extract_text(data, filename=name), False
    except NotImplementedError:
        return (data.decode("utf-8", errors="replace") if name.lower().endswith(".txt")
                else f"Demo content extracted from {name}"), True


def _redact(text: str) -> tuple[str, bool]:
    try:
        from utils.pii import redact_pii
    except (ImportError, ModuleNotFoundError):
        return text, True
    try:
        return redact_pii(text), False
    except NotImplementedError:
        return text, True


def _mock_generate(text: str, slide_count: int, qa_count: int) -> dict[str, Any]:
    topic = text.strip().splitlines()[0][:80] if text.strip() else "Generated presentation"
    deck = {
        "title": topic,
        "subtitle": "A focused proposal tailored to the selected audience",
        "slides": [
            {
                "title": "Problem and context",
                "bullets": [
                    f"{topic} addresses a clear operational or user need",
                    "Current approaches require avoidable time and manual effort",
                    "A focused first release can demonstrate value quickly",
                ],
                "speaker_notes": "Describe the current situation and the people most affected by it.",
                "visual_suggestion": "A before-and-after comparison of the current and proposed experience",
            },
            {
                "title": "User needs",
                "bullets": [
                    "A simple experience that requires little training",
                    "Reliable results with clear explanations and human control",
                    "Privacy safeguards throughout the workflow",
                ],
                "speaker_notes": "Connect each need to a specific user or stakeholder.",
                "visual_suggestion": "Three user profiles with one priority highlighted for each",
            },
            {
                "title": "Proposed solution",
                "bullets": [
                    f"A practical application centred on {topic}",
                    "Guided inputs produce structured, actionable outputs",
                    "Users can review and refine results before using them",
                ],
                "speaker_notes": "Explain the solution in one sentence before discussing features.",
                "visual_suggestion": "Product workflow from user input to reviewed output",
            },
            {
                "title": "How it works",
                "bullets": [
                    "Capture the user's goal and relevant source material",
                    "Process the request with validation and privacy controls",
                    "Return an editable result and preserve useful history",
                ],
                "speaker_notes": "Walk through one realistic example from beginning to end.",
                "visual_suggestion": "Four-step horizontal process diagram",
            },
            {
                "title": "Implementation plan",
                "bullets": [
                    "Begin with one high-value workflow and measurable baseline",
                    "Run a controlled pilot and collect user feedback",
                    "Improve the product before expanding adoption",
                ],
                "speaker_notes": "Position the pilot as a low-risk way to validate the proposal.",
                "visual_suggestion": "Phased roadmap with pilot, evaluation, and expansion",
            },
            {
                "title": "Expected impact",
                "bullets": [
                    "Less time spent on repetitive work",
                    "More consistent and useful outputs",
                    "Clear evidence for the next investment decision",
                ],
                "speaker_notes": "Use measured pilot outcomes instead of unsupported projections.",
                "visual_suggestion": "Three large outcome metrics with baseline placeholders",
            },
        ],
        "qa": copy.deepcopy(MOCK_DECK["qa"]),
    }
    while len(deck["slides"]) < slide_count:
        number = len(deck["slides"]) + 1
        deck["slides"].append({
            "title": f"Strategic Insight {number}",
            "bullets": ["Connect to a measurable user need", "Start with a focused pilot", "Track adoption, quality, cost, and impact"],
            "speaker_notes": "Use one concrete example.",
            "visual_suggestion": "Metric cards or a before-and-after comparison",
        })
    deck["slides"], deck["qa"] = deck["slides"][:slide_count], deck["qa"][:qa_count]
    return deck


def _generate(text: str, settings: dict[str, Any]) -> tuple[Any, bool]:
    try:
        from llm.generator import generate_presentation
    except (ImportError, ModuleNotFoundError):
        return _mock_generate(text, settings["slide_count"], settings["qa_count"]), True
    try:
        try:
            return generate_presentation(source_text=text, **settings), False
        except TypeError:
            return generate_presentation(text, **settings), False
    except NotImplementedError:
        return _mock_generate(text, settings["slide_count"], settings["qa_count"]), True


def _save(deck: Any, settings: dict[str, Any]) -> bool:
    try:
        from db.crud import save_presentation
    except (ImportError, ModuleNotFoundError):
        _DEMO_HISTORY.insert(0, {
            "id": len(_DEMO_HISTORY) + 1,
            "title": _as_dict(deck).get("title", "Untitled presentation"),
            "audience": settings["audience"], "tone": settings["tone"],
            "created_at": "Current session", "deck": copy.deepcopy(deck),
        })
        return True
    try:
        try:
            save_presentation(deck=deck, **settings)
        except TypeError:
            save_presentation(deck, settings)
    except NotImplementedError:
        _DEMO_HISTORY.insert(0, {
            "id": len(_DEMO_HISTORY) + 1,
            "title": _as_dict(deck).get("title", "Untitled presentation"),
            "audience": settings["audience"], "tone": settings["tone"],
            "created_at": "Current session", "deck": copy.deepcopy(deck),
        })
        return True
    return False


def build_deck(*, source_type: str, idea: str, uploaded_bytes: bytes | None,
               uploaded_name: str | None, audience: str, tone: str,
               duration_minutes: int, slide_count: int, qa_count: int) -> tuple[Any, bool]:
    settings = {"audience": audience, "tone": tone, "duration_minutes": duration_minutes,
                "slide_count": slide_count, "qa_count": qa_count}
    demo = False
    if source_type == "document":
        if not uploaded_bytes or not uploaded_name:
            raise ValueError("Please upload a source document.")
        text, fallback = _parse_document(uploaded_bytes, uploaded_name)
        demo |= fallback
    else:
        text = idea.strip()
    if not text:
        raise ValueError("Please provide some source content.")
    text, fallback = _redact(text); demo |= fallback
    deck, fallback = _generate(text, settings); demo |= fallback
    demo |= _save(deck, settings)
    LOGGER.info("Presentation pipeline completed")
    return deck, demo


def regenerate_deck_slide(*, deck: Any, slide_index: int,
                          settings: dict[str, Any]) -> tuple[Any, bool]:
    slides = _as_dict(deck).get("slides", [])
    if not 0 <= slide_index < len(slides):
        raise IndexError("Slide index is out of range.")
    try:
        from llm.generator import regenerate_slide
    except (ImportError, ModuleNotFoundError):
        slides[slide_index]["bullets"] = list(reversed(slides[slide_index].get("bullets", [])))
        slides[slide_index]["speaker_notes"] = (slides[slide_index].get("speaker_notes", "") + " Emphasize the audience outcome.").strip()
        return deck, True
    try:
        replacement = regenerate_slide(slide=slides[slide_index], slide_index=slide_index,
                                       deck=deck, **settings)
    except NotImplementedError:
        slides[slide_index]["bullets"] = list(reversed(slides[slide_index].get("bullets", [])))
        return deck, True
    if isinstance(deck, dict):
        deck["slides"][slide_index] = replacement
    else:
        deck.slides[slide_index] = replacement
    return deck, False


def get_history_items() -> list[Any]:
    try:
        from db.crud import get_history
    except (ImportError, ModuleNotFoundError):
        return _DEMO_HISTORY
    try:
        return get_history()
    except NotImplementedError:
        return _DEMO_HISTORY


def load_history_deck(presentation_id: Any) -> Any:
    try:
        from db.crud import get_presentation
    except (ImportError, ModuleNotFoundError):
        for record in _DEMO_HISTORY:
            if str(record["id"]) == str(presentation_id):
                return copy.deepcopy(record["deck"])
        raise KeyError("Presentation not found in this session.")
    try:
        return get_presentation(presentation_id)
    except NotImplementedError:
        for record in _DEMO_HISTORY:
            if str(record["id"]) == str(presentation_id):
                return copy.deepcopy(record["deck"])
        raise KeyError("Presentation not found in this session.")


def export_deck(deck: Any) -> tuple[bytes, bool]:
    try:
        from utils.exporter import export_pptx
    except (ImportError, ModuleNotFoundError):
        return _fallback_export(deck), True
    try:
        result = export_pptx(deck)
    except NotImplementedError:
        return _fallback_export(deck), True
    if isinstance(result, io.BytesIO):
        return result.getvalue(), False
    if isinstance(result, bytes):
        return result, False
    if isinstance(result, (str, Path)):
        return Path(result).read_bytes(), False
    raise TypeError("export_pptx() must return bytes, BytesIO, or a file path.")


def _fallback_export(deck: Any) -> bytes:
    data = _as_dict(deck)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    navy = RGBColor(15, 23, 42)
    indigo = RGBColor(79, 70, 229)
    violet = RGBColor(124, 58, 237)
    pink = RGBColor(219, 39, 119)
    white = RGBColor(255, 255, 255)
    ink = RGBColor(30, 41, 59)
    muted = RGBColor(100, 116, 139)
    pale = RGBColor(238, 242, 255)
    font = "Aptos"

    def add_text(slide, text, left, top, width, height, size, color,
                 bold=False, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.clear(); frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.text = str(text); paragraph.alignment = align
        paragraph.font.name = font; paragraph.font.size = Pt(size)
        paragraph.font.bold = bold; paragraph.font.color.rgb = color
        return box

    def add_footer(slide, number):
        add_text(slide, "PITCHPILOT AI", 0.68, 7.05, 2.1, 0.22, 8, muted, True)
        add_text(slide, f"{number:02d}", 12.05, 7.0, 0.6, 0.25, 9, muted, True, PP_ALIGN.RIGHT)

    # Minimal branded cover.
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_slide.background.fill.solid(); title_slide.background.fill.fore_color.rgb = navy
    accent = title_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(.18), Inches(7.5))
    accent.fill.solid(); accent.fill.fore_color.rgb = pink; accent.line.fill.background()
    add_text(title_slide, "PITCHPILOT AI", .85, .72, 2.6, .35, 12, RGBColor(196, 181, 253), True)
    add_text(title_slide, data.get("title", "Generated presentation"), .85, 2.05, 11.2, 1.7, 34, white, True)
    add_text(title_slide, data.get("subtitle", "Created with PitchPilot AI"), .88, 4.05, 9.9, .75, 18, RGBColor(203, 213, 225))
    add_text(title_slide, "AI-GENERATED PRESENTATION", .88, 6.55, 3.6, .3, 9, RGBColor(148, 163, 184), True)

    for slide_number, slide_data in enumerate(data.get("slides", []), start=2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = white
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(.18), Inches(7.5))
        band.fill.solid(); band.fill.fore_color.rgb = indigo if slide_number % 2 == 0 else violet
        band.line.fill.background()
        add_text(slide, slide_data.get("title", "Untitled slide"), .72, .55, 8.2, .6, 27, ink, True)
        add_text(slide, "KEY POINTS", .75, 1.55, 2, .3, 10, indigo, True)

        bullets = slide_data.get("bullets", [])[:5]
        for bullet_index, bullet in enumerate(bullets):
            y = 2.0 + bullet_index * 1.05
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(.78), Inches(y + .08), Inches(.18), Inches(.18))
            dot.fill.solid(); dot.fill.fore_color.rgb = pink
            dot.line.fill.background()
            add_text(slide, bullet, 1.15, y, 7.15, .72, 18, ink)

        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.0), Inches(1.48), Inches(3.55), Inches(4.95))
        panel.fill.solid(); panel.fill.fore_color.rgb = pale
        panel.line.color.rgb = RGBColor(224, 231, 255)
        add_text(slide, "VISUAL DIRECTION", 9.38, 1.9, 2.8, .3, 10, violet, True)
        add_text(slide, slide_data.get("visual_suggestion", "Supporting visual"), 9.38, 2.45, 2.75, 1.75, 18, ink, True)
        add_text(slide, "Speaker note", 9.38, 4.8, 2.5, .3, 10, muted, True)
        note_preview = slide_data.get("speaker_notes", "")
        add_text(slide, note_preview, 9.38, 5.18, 2.72, .86, 11, muted)
        slide.notes_slide.notes_text_frame.text = note_preview
        add_footer(slide, slide_number)

    qa = data.get("qa", [])
    if qa:
        number = len(presentation.slides) + 1
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = navy
        add_text(slide, "Questions to expect", .8, .65, 8, .65, 28, white, True)
        for index, item in enumerate(qa[:4]):
            y = 1.65 + index * 1.25
            add_text(slide, f"{index + 1:02d}", .85, y, .55, .35, 13, RGBColor(196, 181, 253), True)
            add_text(slide, item.get("question", "Question"), 1.55, y - .03, 10.5, .4, 17, white, True)
            add_text(slide, item.get("answer", ""), 1.55, y + .43, 10.25, .55, 12, RGBColor(203, 213, 225))
        add_footer(slide, number)

    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()
