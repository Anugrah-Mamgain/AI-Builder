from pathlib import Path

from docx import Document

from utils.parser import (
    EmptyFileError,
    FileTooLargeError,
    NoExtractableTextError,
    UnsupportedFileTypeError,
    extract_text,
)


def test_txt_parser():
    result = extract_text(b"hello world", "test.txt")

    assert result == "hello world"


def test_empty_file_rejected():
    try:
        extract_text(b"", "empty.txt")
        assert False, "Expected EmptyFileError"
    except EmptyFileError:
        pass


def test_large_file_rejected():
    data = b"x" * (5 * 1024 * 1024 + 1)

    try:
        extract_text(data, "large.txt")
        assert False, "Expected FileTooLargeError"
    except FileTooLargeError:
        pass


def test_unsupported_type_rejected():
    try:
        extract_text(b"hello", "malware.exe")
        assert False, "Expected UnsupportedFileTypeError"
    except UnsupportedFileTypeError:
        pass


def test_docx_parser(tmp_path):
    docx_path = tmp_path / "test.docx"

    document = Document()
    document.add_paragraph("AI Builder DOCX test")
    document.add_paragraph("Email: secret@example.com")
    document.save(docx_path)

    result = extract_text(
        docx_path.read_bytes(),
        docx_path.name,
    )

    assert "AI Builder DOCX test" in result
    assert "secret@example.com" in result


def test_pdf_parser():
    pdf_path = Path("sample_inputs/demo_doc.pdf")

    result = extract_text(
        pdf_path.read_bytes(),
        pdf_path.name,
    )

    assert "AI Presentation Builder" in result


def test_invalid_scanned_pdf():
    # Use an image-only PDF placed at sample_inputs/scanned.pdf.
    pdf_path = Path("sample_inputs/scanned.pdf")

    try:
        extract_text(
            pdf_path.read_bytes(),
            pdf_path.name,
        )
        assert False, "Expected NoExtractableTextError"
    except NoExtractableTextError:
        pass