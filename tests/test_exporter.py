from pptx import Presentation as PptxPresentation

from schemas import Presentation
from utils.exporter import export_pptx


def test_export_pptx():
    deck = Presentation(
        title="AI Builder Test",
        audience="Hackathon Judges",
        slides=[
            {
                "slide_no": 1,
                "title": "Problem Statement",
                "key_points": [
                    "Manual presentation creation takes time",
                    "Information is scattered across source documents",
                ],
                "speaker_notes": (
                    "Explain the problem and why automated "
                    "presentation generation is useful."
                ),
                "visual_suggestion": "Document-to-slide pipeline diagram",
            }
        ],
        audience_questions=[
            {
                "question": "How is PII protected?",
                "suggested_answer": (
                    "Sensitive information is redacted before "
                    "the text reaches the LLM."
                ),
            }
        ],
    )

    result = export_pptx(deck)

    assert isinstance(result, bytes)
    assert len(result) > 0

    # A valid PPTX is a ZIP container.
    assert result[:2] == b"PK"

    # Verify PowerPoint can actually open the generated file.
    presentation = PptxPresentation(
        __import__("io").BytesIO(result)
    )

    assert len(presentation.slides) == 3